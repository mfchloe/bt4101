import os
import re

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md"}


def extract_text(path):
    """Return the plain text of a file, based on its extension."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        reader = PdfReader(path)
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        return clean_pdf_text(text)

    if ext == ".docx":
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)

    if ext == ".pptx":
        slides = []
        for i, slide in enumerate(Presentation(path).slides, start=1):
            texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
            slides.append(f"[Slide {i}]\n" + "\n".join(texts))
        return "\n\n".join(slides)

    if ext in (".txt", ".md"):
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read()

    raise ValueError(f"Unsupported file type: {ext}")


def clean_pdf_text(text):
    """Fix characters that some PDFs (e.g. MOE syllabus docs) extract wrongly."""
    text = text.replace("\x96", "–").replace("\x97", "—")  # en/em dashes
    # Bullets drawn with a symbol font come out as "x" at the start of a line
    # or after a number in a table row (e.g. "5 9–10 x All aspects...")
    text = re.sub(r"(?m)(^|(?<=\d ))x ", "• ", text)
    return text

# Can play with chunking strategies later, e.g. chunk by paragraphs or by semantic units
def chunk_text(text, chunk_size=1000, overlap=200):
    """Split text into overlapping chunks of roughly chunk_size characters."""
    text = text.strip()
    chunks = []
    start = 0
    while start < len(text):
        chunks.append(text[start : start + chunk_size])
        start += chunk_size - overlap
    return chunks
